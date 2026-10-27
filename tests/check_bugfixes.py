"""Auto-verificación de los bugs corregidos en la depuración.

Ejecutar:  python tests/check_bugfixes.py
"""
import json
import os
import sys
import tempfile
from unittest import mock

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import core.motor as motor
import core.auditoria as aud
import core.vault as vault
import core.tags as tags
import core.plantillas as plantillas
import core.visor as visor


def check_sql_solo_lectura():
    """El validador bloquea lectura de archivos y sentencias compuestas, pero no literales."""
    assert not motor._es_sql_solo_lectura("SELECT * FROM read_text('data/.vault.key')"), "read_text debe bloquearse"
    assert not motor._es_sql_solo_lectura("SELECT * FROM read_csv_auto('C:/Prototipo/data/users.json')"), "read_csv_auto debe bloquearse"
    assert not motor._es_sql_solo_lectura("SELECT * FROM glob('data/*')"), "glob debe bloquearse"
    assert not motor._es_sql_solo_lectura("SELECT 1; DROP TABLE mantenimientos"), "múltiples sentencias deben bloquearse"
    assert not motor._es_sql_solo_lectura("DELETE FROM mantenimientos"), "escrituras deben bloquearse"
    assert motor._es_sql_solo_lectura("SELECT * FROM mantenimientos WHERE descripcion LIKE '%update%'"), "un literal con palabra prohibida es lectura válida"
    assert motor._es_sql_solo_lectura("SELECT * FROM mantenimientos WHERE estado = 'set'"), "un literal 'set' es lectura válida"
    assert motor._es_sql_solo_lectura("SELECT count(*) FROM mantenimientos")


def check_resaltado_acentos_y_base64():
    """Resalta palabras acentuadas y no toca los data URIs base64."""
    con_acento = "La configuración del balanceador está en revisión."
    res = motor.resaltar_terminos_en_html(con_acento, "configuracion")
    assert "search-highlight" in res and "configuración" in res, "debe resaltar la forma acentuada"

    b64 = "iVBORw0KGgoAAAANSUhEUg=="
    con_img = f"Texto previo ![x](data:image/png;base64,{b64}) texto final"
    no_rompe = motor.resaltar_terminos_en_html(con_img, "data")
    assert f"base64,{b64}" in no_rompe, "el data URI no debe partirse con etiquetas de resaltado"


def check_fragmento_indice_original():
    """El índice del fragmento se calcula sobre el texto original (con emojis y acentos)."""
    contenido = "😀"*50 + "\nConfiguración del cluster: passo normal."
    frag = motor.extraer_fragmento_relevante(contenido, "configuracion", max_chars=200)
    assert "Configuración" in frag, f"el fragmento debe contener el término real: {frag!r}"


def check_visor_resaltado_base64():
    """El buscador de Zen no inyecta <mark> dentro del base64 ni cuenta coincidencias ficticias."""
    b64 = "QUJDREVGR0g="  # contiene "DE" y "F"
    html_in = f"<img src='data:image/png;base64,{b64}'/> palabra data y data"
    res, cnt = visor.resaltar_termino_en_html(html_in, "data")
    assert f"base64,{b64}'" in res, "el src base64 no debe romperse"
    assert cnt == 2, f"las coincidencias deben ser solo texto visible, no el binario (cnt={cnt})"


def check_excel_version_historica():
    """El snapshot de una versión Excel entrega los bytes de esa versión (no los
    actuales) y conserva las demás hojas del libro."""
    import io
    from openpyxl import Workbook, load_workbook
    import pandas as pd

    with tempfile.TemporaryDirectory() as td:
        docs = os.path.join(td, "docs")
        os.makedirs(docs)
        h_orig, d_orig, l_orig = aud.HISTORY_DIR, aud.DOCS_DIR, aud.AUDIT_LOG_PATH
        try:
            aud.DOCS_DIR = docs
            aud.HISTORY_DIR = os.path.join(td, "history")
            aud.AUDIT_LOG_PATH = os.path.join(td, "audit_log.json")
            aud._obtener_historial_versiones_cached.cache_clear()

            libro = os.path.join(docs, "libro.xlsx")
            wb = Workbook()
            ws = wb.active
            ws.title = "Datos"
            ws.append(["servidor", "estado"])
            ws.append(["BALANCER001", "Operativo"])
            wb.create_sheet("Resumen").append(["total", "1"])
            wb.save(libro)

            aud.inicializar_version_inicial_si_no_existe("libro.xlsx", "")
            v = aud.guardar_nueva_version_excel(
                "libro.xlsx", "Datos",
                pd.DataFrame([["BALANCER001", "En revision"]], columns=["servidor", "estado"]),
                autor="QA", comentario="prueba", doc_store={},
            )
            assert v == 2, f"la edición debe crear v2 (v={v})"

            hist = aud.obtener_historial_versiones("libro.xlsx")
            bytes_v1 = aud.obtener_bytes_snapshot("libro.xlsx", hist[0]["archivo_excel_snapshot"])
            with open(libro, "rb") as f:
                assert bytes_v1 and bytes_v1 != f.read(), "el snapshot histórico no debe ser el archivo actual"

            wb_v1 = load_workbook(io.BytesIO(bytes_v1))
            assert wb_v1.sheetnames == ["Datos", "Resumen"], "el snapshot debe conservar las hojas"
            filas_v1 = [[c.value for c in r] for r in wb_v1["Datos"].iter_rows()]
            assert ["BALANCER001", "Operativo"] in filas_v1, f"v1 debe conservar el estado original ({filas_v1})"
        finally:
            aud.HISTORY_DIR, aud.DOCS_DIR, aud.AUDIT_LOG_PATH = h_orig, d_orig, l_orig
            aud._obtener_historial_versiones_cached.cache_clear()


def check_vault_no_destruye_clave():
    """Un fallo de lectura de .vault.key no sobrescribe la clave; sin clave previa sí genera."""
    with tempfile.TemporaryDirectory() as td:
        key_path = os.path.join(td, ".vault.key")
        vault_file = os.path.join(td, ".vault.enc")
        with open(key_path, "wb") as f:
            f.write(b"clave-corrupta-no-fernet")
        with mock.patch.object(vault, "VAULT_KEY_PATH", key_path), \
             mock.patch.object(vault, "VAULT_FILE_PATH", vault_file), \
             mock.patch.object(vault, "ES_PRODUCCION", False), \
             mock.patch.dict(os.environ, {}, clear=False):
            os.environ.pop("VAULT_MASTER_KEY", None)
            try:
                vault._claves_candidatas()
                raise AssertionError("una clave ilegible debe lanzar ErrorBoveda")
            except vault.ErrorBoveda:
                pass
            with open(key_path, "rb") as f:
                assert f.read() == b"clave-corrupta-no-fernet", "la clave ilegible no debe sobrescribirse"

            os.remove(key_path)
            k = vault._claves_candidatas()
            assert os.path.exists(key_path), "sin clave previa debe generarse y persistirse"
            with open(key_path, "rb") as f:
                assert f.read() == k[0]


def check_tags_cuarentena_no_destruye_backup():
    """Un JSON ilegible se cuarentena sin pisar el .corrupto anterior; un error transitorio no renombra."""
    with tempfile.TemporaryDirectory() as td:
        cat_path = os.path.join(td, "categorias.json")
        with mock.patch.object(tags, "CATEGORIAS_PATH", cat_path):
            with open(cat_path, "w", encoding="utf-8") as f:
                f.write("BACKUP-ANTERIOR")
            with open(cat_path + ".corrupto", "w", encoding="utf-8") as f:
                f.write("backup-previo")

            assert tags.cargar_datos_categorias() == {"categorias": [], "documentos": {}}
            with open(cat_path + ".corrupto", "r", encoding="utf-8") as f:
                assert f.read() == "backup-previo", "el .corrupto anterior no debe pisarse"

            # JSON inválido real: como ya existe .corrupto, el nuevo corrupto
            # queda en su sitio (nunca se pisa el único backup recuperable)
            with open(cat_path, "w", encoding="utf-8") as f:
                f.write('{"categorias": [truncado')
            data = tags.cargar_datos_categorias()
            assert data == {"categorias": [], "documentos": {}}
            with open(cat_path + ".corrupto", "r", encoding="utf-8") as f:
                assert f.read() == "backup-previo", "el .corrupto no debe pisarse jamás"
            assert os.path.exists(cat_path), "el archivo corrupto nuevo no debe eliminarse"

            # Fallo de lectura transitorio (OSError simulado): no renombra nada
            with mock.patch("builtins.open", side_effect=OSError("bloqueado")):
                assert tags.cargar_datos_categorias() == {"categorias": [], "documentos": {}}
            with open(cat_path + ".corrupto", "r", encoding="utf-8") as f:
                assert f.read() == "backup-previo", "un OSError transitorio no debe tocar el backup"


def check_plantillas_tipo_raiz():
    """Un plantillas_custom.json con tipo raíz inválido no crashea la pestaña."""
    with tempfile.TemporaryDirectory() as td:
        p = os.path.join(td, "plantillas_custom.json")
        with open(p, "w", encoding="utf-8") as f:
            f.write("[]")
        with mock.patch.object(plantillas, "PLANTILLAS_CUSTOM_PATH", p):
            assert plantillas.cargar_plantillas_personalizadas() == {}
            assert isinstance(plantillas.obtener_todos_los_tipos_plantillas(), list)


if __name__ == "__main__":
    check_sql_solo_lectura()
    check_resaltado_acentos_y_base64()
    check_fragmento_indice_original()
    check_visor_resaltado_base64()
    check_excel_version_historica()
    check_vault_no_destruye_clave()
    check_tags_cuarentena_no_destruye_backup()
    check_plantillas_tipo_raiz()
    print("[OK] check_bugfixes")
