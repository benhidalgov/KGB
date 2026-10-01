import base64
import os
import re
import logging
import streamlit as st
import pandas as pd
from core.auditoria import cargar_hoja_excel_dataframe, guardar_nueva_version, obtener_nombres_hojas_excel

logger = logging.getLogger("infra_copilot.visor")
from core.procesador import IMAGE_EXTENSIONS, preparar_markdown_con_imagenes, normalizar_titulo_display
from core.auth import tiene_permiso
from core.configuracion import DOCS_DIR

MIME_MAP = {
    ".pdf": "application/pdf",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xls": "application/vnd.ms-excel",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
    ".svg": "image/svg+xml", ".webp": "image/webp", ".txt": "text/plain", ".csv": "text/csv"
}


def extraer_caption_diagrama(md_content: str, default_name: str) -> str:
    """Extrae el texto del pie de imagen (caption) definido en la ficha Markdown."""
    if md_content:
        m = re.search(r'\*\*Pie de Imagen\s*(?:\(Caption\))?:\*\*\s*(.+)', md_content, re.IGNORECASE)
        if m and m.group(1).strip():
            return m.group(1).strip()
    return f"Diagrama de Arquitectura: {default_name}"


def actualizar_caption_en_markdown(md_content: str, nuevo_caption: str) -> str:
    """Actualiza o inserta el campo de pie de imagen (caption) en el contenido Markdown."""
    cap_str = nuevo_caption.strip()
    if re.search(r'(\*\*Pie de Imagen\s*(?:\(Caption\))?:\*\*\s*).+', md_content, re.IGNORECASE):
        return re.sub(r'(\*\*Pie de Imagen\s*(?:\(Caption\))?:\*\*\s*).+', lambda m: m.group(1) + cap_str, md_content)
    m_bin = re.search(r'(\* \*\*Archivo Binario:\*\* `[^`]+`\n)', md_content)
    if m_bin:
        return re.sub(r'(\* \*\*Archivo Binario:\*\* `[^`]+`\n)', lambda m: m.group(1) + f"* **Pie de Imagen (Caption):** {cap_str}\n", md_content)
    return f"* **Pie de Imagen (Caption):** {cap_str}\n\n" + md_content


def mostrar_pdf_embebido(pdf_path: str, height: int = 550, key_suffix: str = ""):
    """Renderiza un visor nativo de PDF embebido mediante un iframe Base64 o boton de descarga para archivos pesados."""
    try:
        size_mb = (os.path.getsize(pdf_path) if os.path.exists(pdf_path) else 0) / (1024 * 1024)
        fname = os.path.basename(pdf_path)
        with open(pdf_path, "rb") as f:
            pdf_bytes = f.read()

        if size_mb > 2.5:
            st.markdown(f"""
            <div class="bento-card" style="margin-bottom: 12px;">
                <div style="font-weight: 600; font-size: 0.88rem; color: var(--text-primary); margin-bottom: 4px;">Documento PDF ({size_mb:.1f} MB)</div>
                <div style="font-size: 0.78rem; color: var(--text-secondary); margin-bottom: 10px; line-height: 1.4;">Descárgalo para verlo sin recargar el navegador.</div>
            </div>
            """, unsafe_allow_html=True)
            st.download_button(label=f"Descargar PDF Original ({fname})", data=pdf_bytes, file_name=fname, mime="application/pdf", width="stretch", key=f"dl_heavy_pdf_{fname}_{key_suffix}")
            return

        b64 = base64.b64encode(pdf_bytes).decode("utf-8")
        st.markdown(f'<iframe src="data:application/pdf;base64,{b64}#toolbar=1&navpanes=0" width="100%" height="{height}px" type="application/pdf" style="border:1px solid var(--border-subtle); border-radius:6px; background-color:var(--bg-surface);"></iframe>', unsafe_allow_html=True)
    except Exception as e:
        st.error(f"No se pudo mostrar el PDF: {str(e)}")


@st.cache_data(show_spinner=False)
def cargar_docx_a_html(filepath: str, mtime: float) -> str:
    """Convierte un documento .docx a HTML semántico enriquecido utilizando mammoth."""
    try:
        import mammoth

        def _convertir_imagen(image):
            with image.open() as img_bytes:
                enc = base64.b64encode(img_bytes.read()).decode("ascii")
            return {
                "src": f"data:{image.content_type};base64,{enc}",
                "style": "max-width: 100%; height: auto; border-radius: 6px; margin: 12px 0; box-shadow: 0 1px 3px rgba(0,0,0,0.08);"
            }

        with open(filepath, "rb") as docx_file:
            result = mammoth.convert_to_html(
                docx_file,
                convert_image=mammoth.images.img_element(_convertir_imagen)
            )
            return result.value or ""
    except Exception as e:
        logger.warning(f"[Visor] Error al convertir {filepath} con mammoth: {e}")
        return ""


@st.cache_data(show_spinner=False)
def cargar_pptx_a_slides(filepath: str, mtime: float) -> list:
    """Extrae diapositivas estructuradas de un archivo .pptx con python-pptx."""
    try:
        from pptx import Presentation
        prs = Presentation(filepath)
        slides_data = []
        for idx, slide in enumerate(prs.slides, 1):
            title = ""
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    txt = shape.text_frame.text.strip()
                    if not txt:
                        continue
                    if not title and (shape == slide.shapes.title or idx == 1):
                        title = txt
                    else:
                        texts.append(txt)
            slides_data.append({
                "numero": idx,
                "titulo": title or f"Diapositiva {idx}",
                "contenido": texts
            })
        return slides_data
    except Exception as e:
        logger.warning(f"[Visor] Error al leer {filepath} con python-pptx: {e}")
        return []


def renderizar_diagrama_limpio(ruta_original: str, doc_name: str, md_content: str, ultima_version: int = 1, ultimo_editor: str = "Técnico Responsable", ultimo_timestamp: str = "N/A", key_suffix: str = ""):
    """Renderiza de forma limpia y exclusiva el diagrama/imagen con su caption y formulario de edicion auditada."""
    fname = os.path.basename(ruta_original)
    size_kb = os.path.getsize(ruta_original) / 1024
    ext = os.path.splitext(ruta_original)[1].lower()
    caption_actual = extraer_caption_diagrama(md_content, os.path.splitext(fname)[0])

    st.markdown(f"""
    <div class="bento-card" style="padding: 10px 14px; margin-bottom: 12px; font-size: 0.84rem; display: flex; justify-content: space-between; align-items: center; flex-wrap: wrap; gap: 8px;">
        <div><b>Archivo:</b> <code style="color: var(--text-primary);">{fname}</code></div>
        <div><b>Formato:</b> <span class="badge-ok">{ext.upper().replace('.', '')}</span></div>
        <div><b>Tamaño:</b> <code>{size_kb:.1f} KB</code></div>
        <div><b>Versión:</b> <span class="badge-ok">v{ultima_version}</span></div>
        <div><b>Último Editor:</b> <span style="color: var(--text-secondary);">{ultimo_editor}</span></div>
    </div>
    """, unsafe_allow_html=True)

    with st.container(border=True):
        st.image(ruta_original, caption=caption_actual, width="stretch")

    st.markdown("---")
    st.markdown("##### Editar Descripción del Diagrama")
    st.caption("Cambia la descripción del diagrama. Guarda tu nombre para dejar registro.")

    col_e1, col_e2 = st.columns([1.2, 2])
    with col_e1:
        autor_caption = st.text_input("Tu nombre (*)", key=f"input_author_caption_{doc_name}_{key_suffix}")
        motivo_caption = st.text_input("Motivo", key=f"input_motive_caption_{doc_name}_{key_suffix}")
    with col_e2:
        nuevo_caption_input = st.text_area("Descripción del Diagrama (*)", value=caption_actual, height=108, key=f"textarea_caption_{doc_name}_{key_suffix}")

    col_btn_save, col_btn_info = st.columns([2, 3])
    with col_btn_save:
        if st.button(f"Guardar Descripción (v{ultima_version + 1})", type="primary", width="stretch", key=f"btn_save_caption_{doc_name}_{key_suffix}"):
            if not tiene_permiso("puede_editar_docs"):
                st.error("Tu rol no permite editar descripciones.")
            elif not autor_caption or not autor_caption.strip():
                st.error("Escribe tu nombre.")
            elif not nuevo_caption_input or not nuevo_caption_input.strip():
                st.error("La descripción no puede estar vacía.")
            else:
                md_actualizado = actualizar_caption_en_markdown(md_content, nuevo_caption_input.strip())
                doc_path = os.path.join(DOCS_DIR, doc_name)
                with open(doc_path, "w", encoding="utf-8") as f:
                    f.write(md_actualizado)

                if "doc_store" in st.session_state:
                    st.session_state.doc_store[doc_name] = md_actualizado

                nueva_v = guardar_nueva_version(
                    doc_name=doc_name,
                    nuevo_contenido=md_actualizado,
                    autor=autor_caption.strip(),
                    comentario=motivo_caption.strip() or f"Actualización de pie de imagen: '{nuevo_caption_input.strip()}'",
                    doc_store=st.session_state.get("doc_store")
                )
                st.toast(f"Descripción actualizada (v{nueva_v})")
                st.success(f"Versión v{nueva_v} guardada. Editor: {autor_caption.strip()}.")
                st.rerun()

    with col_btn_info:
        st.caption(f"*Al guardar se creará la versión **v{ultima_version + 1}** y quedará en el historial.*")

    st.markdown("---")
    with open(ruta_original, "rb") as f_img:
        bytes_img = f_img.read()
    mime_type = MIME_MAP.get(ext, "image/png")
    st.download_button(label=f"Descargar Imagen Original ({fname})", data=bytes_img, file_name=fname, mime=mime_type, width="stretch", key=f"dl_btn_diag_direct_{fname}_{key_suffix}")

    with st.expander("Ver Texto Interno (Markdown)", expanded=False):
        st.code(md_content, language="markdown")


def renderizar_original_adaptativo(ruta_original: str, doc_name: str, md_content: str = "", height: int = 480, key_suffix: str = ""):
    """Renderiza el documento fuente original de forma adaptativa según su tipo de formato binario."""
    if not ruta_original or not os.path.exists(ruta_original):
        st.info("Este documento no tiene un archivo original adjunto.")
        return

    ext = os.path.splitext(ruta_original)[1].lower()
    fname = os.path.basename(ruta_original)
    size_kb = os.path.getsize(ruta_original) / 1024

    st.markdown(f"""
    <div class="visor-source-meta-bar">
        <div class="visor-source-meta-title"><b>Archivo Fuente:</b> <code>{fname}</code></div>
        <div class="visor-source-meta-tags">
            <span class="badge-tag">{ext.upper().replace('.', '')}</span>
            <span class="badge-tag">{size_kb:.1f} KB</span>
        </div>
    </div>
    """, unsafe_allow_html=True)

    if ext in IMAGE_EXTENSIONS:
        st.image(ruta_original, caption=extraer_caption_diagrama(md_content, fname), width="stretch")
    elif ext in (".xlsx", ".xls"):
        mtime = os.path.getmtime(ruta_original) if os.path.exists(ruta_original) else 0.0
        sheets = obtener_nombres_hojas_excel(ruta_original, mtime)
        hoja_sel = st.selectbox("Seleccionar Hoja de Trabajo:", sheets, key=f"visor_orig_sheet_{fname}_{key_suffix}")
        df_hoja = cargar_hoja_excel_dataframe(ruta_original, hoja_sel, mtime)
        st.dataframe(df_hoja, width="stretch", height=height)
    elif ext == ".pdf":
        mostrar_pdf_embebido(ruta_original, height=height, key_suffix=key_suffix)
    elif ext in (".docx", ".doc"):
        mtime = os.path.getmtime(ruta_original) if os.path.exists(ruta_original) else 0.0
        html_docx = cargar_docx_a_html(ruta_original, mtime) if ext == ".docx" else ""
        if html_docx:
            st.markdown(
                f'<div class="visor-docx-container" style="max-height: {height}px; overflow-y: auto;">'
                f'{html_docx}'
                f'</div>',
                unsafe_allow_html=True
            )
        elif md_content and md_content.strip():
            with st.container(height=height):
                st.markdown(preparar_markdown_con_imagenes(md_content, doc_name=doc_name, ruta_original=ruta_original), unsafe_allow_html=True)
        else:
            st.info(f"No se pudo generar la vista previa del documento {fname}.")
    elif ext in (".pptx", ".ppt"):
        mtime = os.path.getmtime(ruta_original) if os.path.exists(ruta_original) else 0.0
        slides = cargar_pptx_a_slides(ruta_original, mtime) if ext == ".pptx" else []
        if slides:
            with st.container(height=height):
                for s in slides:
                    bullets_html = "".join(f"<li>{t}</li>" for t in s["contenido"]) if s["contenido"] else "<p style='color:var(--text-muted);font-style:italic;'>Sin contenido textual en la diapositiva.</p>"
                    st.markdown(f"""
                    <div class="visor-pptx-slide-card">
                        <div class="visor-pptx-slide-header">
                            <span class="badge-tag">Diapositiva {s['numero']}</span>
                            <span class="visor-pptx-slide-title">{s['titulo']}</span>
                        </div>
                        <div class="visor-pptx-slide-body">
                            <ul style="margin: 8px 0; padding-left: 20px;">
                                {bullets_html}
                            </ul>
                        </div>
                    </div>
                    """, unsafe_allow_html=True)
        elif md_content and md_content.strip():
            with st.container(height=height):
                st.markdown(preparar_markdown_con_imagenes(md_content, doc_name=doc_name, ruta_original=ruta_original), unsafe_allow_html=True)
        else:
            st.info(f"No se pudo generar la vista previa de la presentación {fname}.")
    elif ext in (".txt", ".csv", ".json", ".sql", ".py", ".md"):
        try:
            with open(ruta_original, "r", encoding="utf-8", errors="ignore") as f:
                raw_text = f.read()
            st.code(raw_text[:8000] + ("\n... [Truncado]" if len(raw_text) > 8000 else ""), language=ext.replace(".", ""))
        except Exception as e:
            st.error(f"Error al leer archivo: {str(e)}")

    with open(ruta_original, "rb") as f:
        bytes_orig = f.read()
    st.download_button(label=f"Descargar Archivo ({fname})", data=bytes_orig, file_name=fname, mime=MIME_MAP.get(ext, "application/octet-stream"), width="stretch", key=f"dl_btn_orig_{fname}_{key_suffix}")


def renderizar_codigo_seguro(md_content: str):
    """Renderiza código Markdown protegiendo el DOM si el texto es muy extenso (> 50 KB)."""
    if len(md_content) > 50_000:
        st.info(f"Documento largo ({len(md_content)/1024:.1f} KB). Se muestran los primeros 50 KB.")
        st.code(md_content[:50_000] + "\n\n... [Truncado en visor de código]", language="markdown")
    else:
        st.code(md_content, language="markdown")


def _render_md_tabs(md_content: str, doc_name: str, ruta_original: str | None):
    tab_rendered, tab_source = st.tabs(["Vista Formateada", "Código Markdown"])
    with tab_rendered:
        with st.container(border=True):
            st.markdown(preparar_markdown_con_imagenes(md_content, doc_name=doc_name, ruta_original=ruta_original), unsafe_allow_html=True)
    with tab_source:
        renderizar_codigo_seguro(md_content)


def renderizar_lado_a_lado(doc_name: str, md_content: str, ruta_original: str | None, ultima_version: int = 1, ultimo_editor: str = "Técnico Responsable", ultimo_timestamp: str = "N/A", key_suffix: str = ""):
    """Gestiona la visualización Lado a Lado de documentos técnicos y diagramas con Modo Enfoque inmersivo."""
    es_diag = doc_name.startswith("DIAGRAMA__") or (ruta_original and any(ruta_original.lower().endswith(e) for e in IMAGE_EXTENSIONS))
    if es_diag and ruta_original and os.path.exists(ruta_original):
        renderizar_diagrama_limpio(ruta_original, doc_name, md_content, ultima_version, ultimo_editor, ultimo_timestamp, key_suffix)
        return

    col_sel, col_stat, col_zen = st.columns([2.5, 1.2, 1.3], vertical_alignment="center")
    with col_sel:
        modo_vista = st.segmented_control(
            "Modo de Visualización",
            ["[Lado a Lado]", "[Solo Markdown]", "[Solo Formato Original]"],
            default="[Lado a Lado]",
            label_visibility="collapsed",
            key=f"seg_modo_vista_{doc_name}_{key_suffix}"
        ) or "[Lado a Lado]"
    with col_stat:
        badge = '<span class="badge-ok" style="font-size:0.72rem;padding:2px 7px;">Fuente Disponible</span>' if (ruta_original and os.path.exists(ruta_original)) else '<span class="badge-warn" style="font-size:0.72rem;padding:2px 7px;">Nativo Markdown</span>'
        st.markdown(f'<div style="text-align:center;font-size:0.78rem;opacity:0.9;"><b>Estado:</b> {badge}</div>', unsafe_allow_html=True)
    with col_zen:
        if st.button("Abrir en Modo Lectura", type="primary", width="stretch", key=f"btn_zen_enter_{doc_name}_{key_suffix}", help="Abre el modo de lectura a pantalla completa."):
            st.session_state["zen_studio_activo"] = True
            st.session_state["zen_doc_sel"] = doc_name
            st.rerun()

    alt_visores = 550
    st.markdown("<div style='height:8px;'></div>", unsafe_allow_html=True)

    if modo_vista == "[Lado a Lado]":
        col_md, col_orig = st.columns(2, gap="medium")
        with col_md:
            st.markdown("##### [Texto]")
            _render_md_tabs(md_content, doc_name, ruta_original)
        with col_orig:
            st.markdown("##### [Archivo Original]")
            with st.container(border=True):
                renderizar_original_adaptativo(ruta_original, doc_name, md_content=md_content, height=alt_visores, key_suffix=f"side_{key_suffix}")
    elif modo_vista == "[Solo Markdown]":
        _render_md_tabs(md_content, doc_name, ruta_original)
    else:
        with st.container(border=True):
            renderizar_original_adaptativo(ruta_original, doc_name, md_content=md_content, height=alt_visores, key_suffix=f"full_{key_suffix}")


def extraer_tabla_de_contenidos(md_content: str) -> list[dict]:
    """Analiza el documento Markdown y extrae la jerarquía de títulos (H1-H4)."""
    toc = []
    lineas = md_content.splitlines()
    for idx, linea in enumerate(lineas):
        m = re.match(r'^(#{1,4})\s+(.+)$', linea.strip())
        if m:
            nivel = len(m.group(1))
            titulo = m.group(2).strip()
            titulo_limpio = re.sub(r'[*_`]', '', titulo)
            toc.append({
                "nivel": nivel,
                "titulo": titulo_limpio,
                "linea": idx + 1,
                "id": f"sec_{idx}"
            })
    return toc


def resaltar_termino_en_html(html_o_md: str, termino: str) -> tuple[str, int]:
    """Resalta con etiqueta mark las ocurrencias de una palabra clave."""
    if not termino or not termino.strip():
        return html_o_md, 0
    t_clean = re.escape(termino.strip())
    patron = re.compile(rf'(?i)({t_clean})')
    coincidencias = len(patron.findall(html_o_md))
    if coincidencias > 0:
        res = patron.sub(r'<mark style="background-color: var(--pastel-amber-bg); color: var(--pastel-amber-text); padding: 2px 4px; border-radius: 3px; font-weight: 600;">\1</mark>', html_o_md)
        return res, coincidencias
    return html_o_md, 0


def renderizar_zen_studio(doc_name: str, md_content: str, ruta_original: str | None, u_ver: int = 1, u_edit: str = "Técnico", u_time: str = "N/A"):
    """Renderiza el entorno inmersivo Zen Studio con TOC interactivo, buscador interno, personalización de lectura y visor multimodal."""
    if hasattr(st, "html"):
        st.html("<style>section[data-testid=\"stSidebar\"], header[data-testid=\"stHeader\"] { display: none !important; } [data-testid=\"stMainBlockContainer\"] { max-width: 98vw !important; padding: 0.6rem 1.4rem !important; }</style>")
    else:
        st.markdown("<style>section[data-testid=\"stSidebar\"], header[data-testid=\"stHeader\"] { display: none !important; } [data-testid=\"stMainBlockContainer\"] { max-width: 98vw !important; padding: 0.6rem 1.4rem !important; }</style>", unsafe_allow_html=True)

    col_zt_title, col_zt_search, col_zt_theme, col_zt_exit = st.columns([3.2, 2.4, 2.2, 1.2], vertical_alignment="center")

    with col_zt_title:
        st.markdown(f"""
        <div style="display:flex;align-items:center;gap:8px;padding-top:4px;">
            <span class="badge-ok" style="font-size:0.7rem;padding:2px 8px;font-weight:600;">[MODO LECTURA]</span>
            <span style="font-family:var(--font-serif);font-size:1.2rem;font-weight:500;color:var(--text-primary);">{normalizar_titulo_display(doc_name)}</span>
            <span class="badge-info" style="font-size:0.68rem;">v{u_ver}</span>
        </div>
        """, unsafe_allow_html=True)

    with col_zt_search:
        zen_search_query = st.text_input("Buscar en doc:", label_visibility="collapsed", key="zen_search_in_doc")

    with col_zt_theme:
        col_zt_th1, col_zt_th2 = st.columns(2)
        with col_zt_th1:
            tema_lectura = st.selectbox("Tema", ["Cálido", "Sepia", "Grafito"], label_visibility="collapsed", key="zen_theme_selector")
        with col_zt_th2:
            tam_fuente = st.selectbox("Tamaño", ["Normal (15px)", "Grande (17px)", "Compacto (13px)"], label_visibility="collapsed", key="zen_font_size_selector")

    with col_zt_exit:
        if st.button("Salir", type="primary", width="stretch", key="btn_exit_zen_studio", help="Vuelve a la consola de operaciones"):
            st.session_state["zen_studio_activo"] = False
            # El router (app.py) decide por nav_seccion_activa; sin esto el
            # siguiente rerun vuelve a entrar en Zen y queda en bucle.
            st.session_state["nav_seccion_activa"] = "Consultas y Búsqueda"
            st.rerun()

    font_size_val = "17px" if "Grande" in tam_fuente else ("13px" if "Compacto" in tam_fuente else "15px")

    paletas_lectura = {
        "Sepia": {
            "bg": "#F5EEDC",
            "fg": "#342A1E",
            "fg_subtle": "#544635",
            "border": "#DECFA8",
            "code_bg": "#EADFCA",
            "card_bg": "#EFE6D2",
        },
        "Grafito": {
            "bg": "#1C1E22",
            "fg": "#EDEDEC",
            "fg_subtle": "#A8A7A0",
            "border": "rgba(255, 255, 255, 0.15)",
            "code_bg": "rgba(255, 255, 255, 0.08)",
            "card_bg": "#24272D",
        },
        "Cálido": {
            "bg": "#FAF7F2",
            "fg": "#1C1B19",
            "fg_subtle": "#484540",
            "border": "#E5E0D5",
            "code_bg": "#EFEBE2",
            "card_bg": "#F4EFE6",
        },
    }
    p_lec = paletas_lectura.get(tema_lectura, paletas_lectura["Cálido"])

    css_zen = f"""
    <style>
    .zen-reader-canvas {{
        background-color: {p_lec['bg']} !important;
        border: 1px solid {p_lec['border']} !important;
        color: {p_lec['fg']} !important;
        font-size: {font_size_val} !important;
        line-height: 1.75 !important;
        border-radius: 8px !important;
        box-shadow: 0 1px 3px rgba(0, 0, 0, 0.04) !important;
    }}
    .zen-reader-canvas p,
    .zen-reader-canvas li,
    .zen-reader-canvas div,
    .zen-reader-canvas span,
    .zen-reader-canvas td,
    .zen-reader-canvas th {{
        color: {p_lec['fg_subtle']} !important;
        font-size: {font_size_val} !important;
        line-height: 1.75 !important;
    }}
    .zen-reader-canvas h1,
    .zen-reader-canvas h2,
    .zen-reader-canvas h3,
    .zen-reader-canvas h4,
    .zen-reader-canvas h5,
    .zen-reader-canvas h6,
    .zen-reader-canvas strong,
    .zen-reader-canvas b {{
        color: {p_lec['fg']} !important;
        font-family: var(--font-serif) !important;
    }}
    .zen-reader-canvas code,
    .zen-reader-canvas pre {{
        background-color: {p_lec['code_bg']} !important;
        color: {p_lec['fg']} !important;
        border: 1px solid {p_lec['border']} !important;
        font-size: 0.88em !important;
    }}
    .zen-reader-canvas blockquote {{
        border-left: 3px solid {p_lec['border']} !important;
        color: {p_lec['fg_subtle']} !important;
        padding-left: 12px !important;
        margin-left: 0 !important;
    }}
    .zen-reader-canvas hr {{
        border: none !important;
        border-top: 1px solid {p_lec['border']} !important;
        margin: 20px 0 !important;
    }}
    .zen-reader-canvas table {{
        border: 1px solid {p_lec['border']} !important;
    }}
    .zen-toc-card {{
        background-color: {p_lec['card_bg']} !important;
        border: 1px solid {p_lec['border']} !important;
    }}
    .zen-toc-item {{
        color: {p_lec['fg_subtle']} !important;
    }}
    .zen-toc-item:hover {{
        background-color: {p_lec['bg']} !important;
        color: {p_lec['fg']} !important;
    }}
    </style>
    """.strip()
    st.markdown(css_zen, unsafe_allow_html=True)

    col_toc, col_canvas = st.columns([1.1, 3.4], gap="medium")

    with col_toc:
        st.markdown("##### [Índice]")
        toc_items = extraer_tabla_de_contenidos(md_content)
        palabras = len(md_content.split())
        minutos_lectura = max(1, palabras // 200)

        if toc_items:
            with st.container(border=True):
                st.caption(f"**{len(toc_items)} secciones**:")
                opciones_seccion = ["Documento Completo"] + [f"{'—' * (it['nivel'] - 1)} {it['titulo']}" for it in toc_items]
                seccion_sel = st.selectbox("Saltar a Sección:", opciones_seccion, key=f"zen_section_jump_sel_{doc_name}")

                toc_html_list = []
                for it in toc_items:
                    indent = (it['nivel'] - 1) * 10
                    badge_h = f"<span class='badge-tag' style='font-size:0.65rem;margin-right:4px;'>H{it['nivel']}</span>"
                    toc_html_list.append(f"<div class='zen-toc-item' style='margin-left:{indent}px;'>{badge_h} {it['titulo']}</div>")
                st.markdown(f"<div class='zen-toc-card'>{''.join(toc_html_list)}</div>", unsafe_allow_html=True)
        else:
            st.info("Este documento no tiene secciones.")
            seccion_sel = "Documento Completo"

        st.markdown(f"""
        <div class="bento-card" style="padding:10px 14px;margin-top:10px;font-size:0.8rem;">
            <div style="font-weight:600;color:var(--text-primary);margin-bottom:6px;">Detalles del Documento</div>
            <div style="display:flex;flex-direction:column;gap:4px;color:var(--text-secondary);">
                <div><b>Palabras:</b> {palabras:,}</div>
                <div><b>Lectura:</b> ~{minutos_lectura} min</div>
                <div><b>Editor:</b> {u_edit}</div>
                <div><b>Actualizado:</b> {u_time}</div>
            </div>
        </div>
        """, unsafe_allow_html=True)

        if ruta_original and os.path.exists(ruta_original):
            ext_orig = os.path.splitext(ruta_original)[1].lower()
            with open(ruta_original, "rb") as f_dl_zen:
                st.download_button(
                    label=f"Descargar Fuente ({ext_orig.upper().replace('.', '')})",
                    data=f_dl_zen.read(),
                    file_name=os.path.basename(ruta_original),
                    mime=MIME_MAP.get(ext_orig, "application/octet-stream"),
                    width="stretch",
                    key=f"zen_dl_btn_{doc_name}"
                )

    with col_canvas:
        texto_a_mostrar = md_content
        if seccion_sel != "Documento Completo":
            # El TOC limpia énfasis (*_`); comparar los encabezados igualmente limpiados.
            titulo_buscado = seccion_sel.lstrip('— ').strip()
            m_sec = None
            for m_h in re.finditer(r'(?m)^#+\s+(.+)$', md_content):
                if re.sub(r'[*_`]', '', m_h.group(1)).strip() == titulo_buscado:
                    resto = md_content[m_h.end():]
                    sig = re.search(r'(?m)^#+\s+', resto)
                    m_sec = md_content[m_h.start():m_h.end() + (sig.start() if sig else len(resto))]
                    break
            if m_sec:
                texto_a_mostrar = m_sec

        cnt_coincidencias = 0
        if zen_search_query.strip():
            texto_preparado = preparar_markdown_con_imagenes(texto_a_mostrar, doc_name=doc_name, ruta_original=ruta_original)
            html_resaltado, cnt_coincidencias = resaltar_termino_en_html(texto_preparado, zen_search_query)
            st.markdown(f"<div style='font-size:0.8rem;margin-bottom:6px;'><span class='badge-ok'>[{cnt_coincidencias} Coincidencias]</span> para '<b>{zen_search_query}</b>'</div>", unsafe_allow_html=True)
        else:
            html_resaltado = preparar_markdown_con_imagenes(texto_a_mostrar, doc_name=doc_name, ruta_original=ruta_original)

        tiene_orig = ruta_original and os.path.exists(ruta_original)
        if tiene_orig:
            tab_z_fmt, tab_z_orig, tab_z_side = st.tabs(["Texto", "Original", "Lado a Lado"])
            with tab_z_fmt:
                st.markdown(f"<div class='zen-reader-canvas'>{html_resaltado}</div>", unsafe_allow_html=True)
            with tab_z_orig:
                with st.container(border=True):
                    renderizar_original_adaptativo(ruta_original, doc_name, md_content=md_content, height=750, key_suffix="zen_orig")
            with tab_z_side:
                col_z_s1, col_z_s2 = st.columns(2, gap="medium")
                with col_z_s1:
                    st.markdown(f"<div class='zen-reader-canvas' style='padding:16px;'>{html_resaltado}</div>", unsafe_allow_html=True)
                with col_z_s2:
                    with st.container(border=True):
                        renderizar_original_adaptativo(ruta_original, doc_name, md_content=md_content, height=750, key_suffix="zen_side")
        else:
            st.markdown(f"<div class='zen-reader-canvas'>{html_resaltado}</div>", unsafe_allow_html=True)
